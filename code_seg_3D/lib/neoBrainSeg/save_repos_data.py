<<<<<<< ours
import numpy as np
import scipy.io as sio
import pathlib
from lib.neoBrainSeg import find_ref_patients_xls_line
import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter
import os
from tools.export_to_3d_dicom import export_to_3d_dicom
from tools.export_rawmhd import export_rawmhd

def save_repos_data(data_sag, data_repos, delta_X_sag, delta_X_seqdyn, idx, main_dir, num_patient, ref, mode_debug):
    """
    Saves processed arrays to .mat files and updates the Excel reference sheet.
    """
    base_path = pathlib.Path(main_dir)
    repos_dir = base_path / 'Pre_traitement_echo_v2' / 'Cropping' / str(ref)
    ref_files_dir = base_path / 'Ref_files'

    repos_dir.mkdir(parents=True, exist_ok=True)
    
    # Sauvegarde des fichiers .mat (SANS processEvents)
    sio.savemat(repos_dir / f"data_repos_{ref}.mat", 
                {'data_repos': data_repos}, do_compression=True)
    
    sio.savemat(repos_dir / f"data_repos_{ref}_sag.mat", 
                {'data_sag': data_sag}, do_compression=True)
    print("dtype", data_repos.dtype)

    # Export Dicom 3D
    data_repos_db = data_repos.astype(float)
    #export_data_mhd_raw_v2.export_data_mhd_raw_v2(data_repos_db, ref, num_patient, 'repos', main_dir, 0)
    sub_dir = os.path.join('Pre_traitement_echo_v2', 'Cropping', str(ref))
    filename = f"data_repos_{ref}"

    # Export vers le DICOM 3D Unique (Utilise les données 0-255)
    # Sécurité max pour éviter la division par zéro
    max_val = np.max(data_repos_db)
    if max_val == 0:
        max_val = 1.0
        print(f"Attention: Volume vide pour {ref}")
    data_uint8 = (data_repos_db.astype(np.float32) / max_val * 255).astype(np.uint8)
    output_dicom_file = os.path.join(main_dir, sub_dir, f"{filename}.dcm")

    export_to_3d_dicom(
        data=data_uint8, 
        output_filepath=output_dicom_file, 
        patient_name=str(num_patient), 
        patient_id=str(ref)
    )

    if mode_debug:  
        # Mise à jour .mat des paramètres
        params_path = ref_files_dir / 'ReconstructionParameters_v2.mat'
        
        try:
            if params_path.exists():
                mat_contents = sio.loadmat(str(params_path))
                params_v2 = mat_contents['params_patients_echo_v2']
            else:
                # Création d'un tableau NumPy rempli de zéros
                # Format : (lignes, colonnes)
                params_v2 = np.zeros((498, 14)) 
                print(f"Nouveau tableau params_v2 créé : {params_v2.shape}")

            # Cette syntaxe [row, col] fonctionne maintenant parfaitement
            params_v2[idx+1, 12] = delta_X_seqdyn
            params_v2[idx+1, 13] = delta_X_sag
            print(params_v2[idx+1])
            
            sio.savemat(str(params_path), {'params_patients_echo_v2': params_v2})

        except Exception as e:
            print(f"Warning: Could not update .mat file: {e}")

    # Mise à jour Excel
    excel_path = ref_files_dir / 'RefPatientsUS.xlsx'
    
    def update_excel_cell(param_name, value, offset=7):
        begin_xls, end_xls, xls_sheet, nb_cell = find_ref_patients_xls_line.find_ref_patients_xls_line(main_dir, param_name)
        
        # On ouvre, modifie et ferme le classeur proprement
        wb = openpyxl.load_workbook(excel_path)
        ws = wb[xls_sheet] if xls_sheet in wb.sheetnames else wb.active
        new_col_letter = get_column_letter(column_index_from_string(begin_xls) - offset)
        cell_address = f"{new_col_letter}{idx + 1 + 2}"
        ws[cell_address] = value
        wb.save(excel_path)
        wb.close() # Important pour libérer le fichier

    update_excel_cell('delta_X_cc', delta_X_sag)
    update_excel_cell('delta_X_seqdyn', delta_X_seqdyn)
=======
import numpy as np
import scipy.io as sio
import pathlib
from lib.neoBrainSeg import find_ref_patients_xls_line
import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter
import os
from tools.export_to_3d_dicom import export_to_3d_dicom
from tools.export_rawmhd import export_rawmhd

def save_repos_data(data_sag, data_repos, delta_X_sag, delta_X_seqdyn, idx, main_dir, num_patient, ref, mode_debug):
    """
    Saves processed arrays to .mat files and updates the Excel reference sheet.
    """
    base_path = pathlib.Path(main_dir)
    repos_dir = base_path / 'Pre_traitement_echo_v2' / 'Cropping' / str(ref)
    ref_files_dir = base_path / 'Ref_files'

    repos_dir.mkdir(parents=True, exist_ok=True)
    
    # Sauvegarde des fichiers .mat (SANS processEvents)
    sio.savemat(repos_dir / f"data_repos_{ref}.mat", 
                {'data_repos': data_repos}, do_compression=True)
    
    sio.savemat(repos_dir / f"data_repos_{ref}_sag.mat", 
                {'data_sag': data_sag}, do_compression=True)
    print("dtype", data_repos.dtype)

    # Export Dicom 3D
    data_repos_db = data_repos.astype(float)
    #export_data_mhd_raw_v2.export_data_mhd_raw_v2(data_repos_db, ref, num_patient, 'repos', main_dir, 0)
    sub_dir = os.path.join('Pre_traitement_echo_v2', 'Cropping', str(ref))
    filename = f"data_repos_{ref}"

    # Export vers le DICOM 3D Unique (Utilise les données 0-255)
    # Sécurité max pour éviter la division par zéro
    max_val = np.max(data_repos_db)
    if max_val == 0:
        max_val = 1.0
        print(f"Attention: Volume vide pour {ref}")
    data_uint8 = (data_repos_db.astype(np.float32) / max_val * 255).astype(np.uint8)
    output_dicom_file = os.path.join(main_dir, sub_dir, f"{filename}.dcm")

    export_to_3d_dicom(
        data=data_uint8, 
        output_filepath=output_dicom_file, 
        patient_name=str(num_patient), 
        patient_id=str(ref)
    )

    if mode_debug:  
        # Mise à jour .mat des paramètres
        params_path = ref_files_dir / 'ReconstructionParameters_v2.mat'
        
        try:
            if params_path.exists():
                mat_contents = sio.loadmat(str(params_path))
                params_v2 = mat_contents['params_patients_echo_v2']
            else:
                # Création d'un tableau NumPy rempli de zéros
                # Format : (lignes, colonnes)
                params_v2 = np.zeros((498, 14)) 
                print(f"Nouveau tableau params_v2 créé : {params_v2.shape}")

            # Cette syntaxe [row, col] fonctionne maintenant parfaitement
            params_v2[idx+1, 12] = delta_X_seqdyn
            params_v2[idx+1, 13] = delta_X_sag
            print(params_v2[idx+1])
            
            sio.savemat(str(params_path), {'params_patients_echo_v2': params_v2})

        except Exception as e:
            print(f"Warning: Could not update .mat file: {e}")

    # Mise à jour Excel
    excel_path = ref_files_dir / 'RefPatientsUS.xlsx'
    
    def update_excel_cell(param_name, value, offset=7):
        begin_xls, end_xls, xls_sheet, nb_cell = find_ref_patients_xls_line.find_ref_patients_xls_line(main_dir, param_name)
        
        # On ouvre, modifie et ferme le classeur proprement
        wb = openpyxl.load_workbook(excel_path)
        ws = wb[xls_sheet] if xls_sheet in wb.sheetnames else wb.active
        new_col_letter = get_column_letter(column_index_from_string(begin_xls) - offset)
        cell_address = f"{new_col_letter}{idx + 1 + 2}"
        ws[cell_address] = value
        wb.save(excel_path)
        wb.close() # Important pour libérer le fichier

    update_excel_cell('delta_X_cc', delta_X_sag)
    update_excel_cell('delta_X_seqdyn', delta_X_seqdyn)
>>>>>>> theirs
